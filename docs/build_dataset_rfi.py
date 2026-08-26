from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(__file__).with_name("Gen_Audius_LLC_QAMUZ_Licensed_Music_Dataset_RFI.docx")

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT = "E8EEF5"
PALE = "F4F6F9"
GRAY = "667085"
WHITE = "FFFFFF"
BLACK = "111827"


def set_font(run, name="Calibri", size=11, bold=False, italic=False, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_table_geometry(table, widths_dxa, indent=120):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths_dxa)))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), str(indent))
    tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(widths_dxa[idx]))
            tcW.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths_dxa):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_font(r, size=9.5, bold=True, color=NAVY)
    set_repeat_table_header(table.rows[0])
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_font(r, size=9.5)
    set_table_geometry(table, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_font(r)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_font(r)
    return p


def add_label_para(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    a = p.add_run(label + ": ")
    set_font(a, bold=True, color=NAVY)
    b = p.add_run(text)
    set_font(b)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25
for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 14, 8),
    ("Heading 2", 13, BLUE, 11, 6),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
]:
    st = styles[name]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for list_name in ["List Bullet", "List Bullet 2", "List Number"]:
    st = styles[list_name]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.25

# Running header/footer
hp = section.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hr = hp.add_run("GEN AUDIUS LLC  |  QAMUZ DATA LICENSING RFI")
set_font(hr, size=8.5, bold=True, color=GRAY)
fp = section.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fr = fp.add_run("Confidential - Vendor evaluation copy")
set_font(fr, size=8, color=GRAY)

# Proposal centerpiece
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(18)
p.paragraph_format.space_after = Pt(6)
r = p.add_run("GEN AUDIUS LLC")
set_font(r, size=12, bold=True, color=GRAY)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(5)
r = p.add_run("REQUEST FOR INFORMATION & QUOTATION")
set_font(r, size=24, bold=True, color=NAVY)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run("Licensed Tropical Latin Music and Singing-Voice Datasets for QAMUZ")
set_font(r, size=14, color=DARK_BLUE)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(24)
r = p.add_run("Pilot acquisition: 50, 100, and 300 hours | Commercial generative-AI training")
set_font(r, size=10.5, bold=True, color=GRAY)

add_table(doc, ["Prepared by", "Contact", "Issued", "Response requested"], [["Gen Audius LLC / QAMUZ", "info@genaudius.com", "August 4, 2026", "Within 14 calendar days"]], [2250, 2730, 1980, 2400])

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(10)
r = p.add_run("Purpose")
set_font(r, size=11, bold=True, color=NAVY)
r = p.add_run("  Gen Audius LLC is evaluating ethically sourced, rights-cleared datasets to train QAMUZ, its flagship brand and commercial music-generation platform. The immediate priority is tropical Latin music; a separate symbolic MIDI program will cover broader genres.")
set_font(r)

doc.add_heading("1. Company and project overview", level=1)
add_label_para(doc, "Legal entity", "Gen Audius LLC")
add_label_para(doc, "Brand and product", "QAMUZ - the flagship public-facing brand and multilingual generative music platform of Gen Audius LLC")
add_label_para(doc, "Current focus", "Long-form music generation, musical structure, singing voices, accompaniment, and cover-art workflows, developed initially with local models.")
add_label_para(doc, "Procurement objective", "Secure for Gen Audius LLC a commercially usable pilot dataset with documented provenance, artist consent, and rights broad enough to train, fine-tune, evaluate, deploy, and improve QAMUZ generative models.")

doc.add_page_break()
doc.add_heading("2. Requested tropical Latin scope", level=1)
add_table(doc, ["Genre family", "Target share", "Desired coverage"], [
    ("Bachata", "35%", "Traditional, modern, urban; male/female leads, choruses, future duets"),
    ("Salsa", "20%", "Romantica, dura, son-based, contemporary arrangements"),
    ("Merengue", "15%", "Traditional, orchestral, mambo and contemporary"),
    ("Cumbia", "10%", "Caribbean, Colombian and modern regional variants"),
    ("Son / bolero / guaracha", "10%", "Foundational tropical harmony, rhythm and phrasing"),
    ("Tropical fusion", "10%", "Latin pop, Afro-Caribbean and modern crossover"),
], [2160, 1260, 5940])
add_bullet(doc, "Please quote three pilot tiers: approximately 50, 100, and 300 usable hours after quality control.")
add_bullet(doc, "Indicate how much content is exclusive, non-exclusive, synthetic, human-performed, previously released, or commissioned.")
add_bullet(doc, "Provide country/region representation and Spanish-language/dialect coverage.")

doc.add_heading("3. Required data package", level=1)
for text in [
    "Full-length masters in lossless WAV (preferred 44.1 or 48 kHz; 24-bit where available).",
    "Stems where available: lead vocal, backing vocals, percussion, bass, guitars, keys, brass, strings and other instruments.",
    "Authorized lyrics and time-aligned lyric segments where available.",
    "Metadata: genre, subgenre, mood, BPM, key, meter, language, territory, instrumentation, vocal configuration and arrangement structure.",
    "Train/validation/test recommendations and duplicate or near-duplicate identifiers.",
    "Checksums, manifest files and stable internal track identifiers.",
    "A small evaluation sample before contract execution.",
]:
    add_bullet(doc, text)

doc.add_heading("4. Singing-voice requirements", level=1)
add_table(doc, ["Category", "Requested information"], [
    ("Voice coverage", "Male, female, backing chorus and duet-capable material; voice type/range tags preferred."),
    ("Consent", "Written performer consent for AI/ML training and commercial generative outputs."),
    ("Identity restrictions", "Disclose restrictions on artist name, likeness, voice cloning, impersonation and publicity rights."),
    ("Technical assets", "Dry/wet vocal stems, isolated phrases, full performances, lyrics, phoneme or word timing if available."),
    ("Language", "Spanish is the priority; identify Caribbean regional pronunciation and any bilingual content."),
], [2160, 7200])

doc.add_heading("5. Rights and licensing requirements", level=1)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(6)
r = p.add_run("The proposal should explicitly confirm whether the license permits each of the following:")
set_font(r, bold=True, color=NAVY)
for text in [
    "Commercial machine-learning training, fine-tuning, validation, evaluation and repeated retraining.",
    "Creation and commercial exploitation of model outputs by Gen Audius LLC and QAMUZ users.",
    "Internal storage, preprocessing, feature extraction, segmentation, augmentation and format conversion.",
    "Deployment of trained models through web applications, APIs, mobile products and enterprise integrations.",
    "Distribution or hosting of trained weights and adapters, including through cloud infrastructure and private customer deployments.",
    "Worldwide use, stated term, renewal terms, termination effects and treatment of models trained before termination.",
    "Coverage of master, composition, lyrics, performer consent, neighboring rights, publicity/voice rights and all necessary sublicenses.",
    "Representations, warranties, provenance records, audit support, takedown process and indemnification.",
]:
    add_bullet(doc, text)

doc.add_heading("6. Pricing questionnaire", level=1)
add_table(doc, ["Pricing item", "Vendor response requested"], [
    ("Pilot tiers", "Price for 50, 100 and 300 usable hours; minimum commitment."),
    ("Asset premiums", "Separate pricing for stems, MIDI, lyrics, annotations and isolated vocals."),
    ("License model", "Perpetual vs. term-limited; exclusive vs. non-exclusive; renewal/escalation."),
    ("Usage scale", "Limits tied to users, revenue, compute, models, products or generated outputs."),
    ("Royalties", "Any recurring fees, output royalties, revenue share or reporting duties."),
    ("Updates", "Pricing and cadence for future catalog refreshes or incremental deliveries."),
    ("Evaluation", "Availability and terms of a small sample or paid proof of concept."),
], [2700, 6660])

doc.add_heading("7. Vendor response checklist", level=1)
for text in [
    "Company and primary licensing contact.",
    "Available hours/tracks by requested genre and voice configuration.",
    "Example manifest and metadata dictionary.",
    "Rights chain and consent documentation summary.",
    "Proposed license scope and a sample agreement.",
    "Pricing for all three pilot sizes and optional assets.",
    "Delivery method, estimated schedule and technical support.",
    "Known exclusions, restrictions and indemnification limits.",
]:
    add_number(doc, text)

doc.add_heading("8. Evaluation criteria", level=1)
add_table(doc, ["Criterion", "Weight", "What Gen Audius LLC will evaluate"], [
    ("Legal clarity and provenance", "30%", "Documented rights, consent, commercial scope and indemnity for Gen Audius LLC"),
    ("Tropical Latin relevance", "25%", "Genre authenticity, regional breadth and arrangement quality"),
    ("Audio and annotation quality", "20%", "Lossless sources, stems, metadata and consistency"),
    ("Voice diversity", "10%", "Male/female leads, choruses, duet material and Spanish coverage"),
    ("Price and scalability", "10%", "Pilot affordability and growth path"),
    ("Delivery and support", "5%", "Speed, documentation and technical responsiveness"),
], [2700, 1080, 5580])

doc.add_heading("9. Initial vendor-specific inquiry messages", level=1)

messages = [
    ("GCX / Rightsify", "team@gcx.co", "Licensed tropical Latin dataset for Gen Audius LLC / QAMUZ", "Gen Audius LLC is evaluating a rights-cleared tropical Latin music dataset for QAMUZ, our commercial generative-AI music brand and platform. Please quote 50-, 100-, and 300-hour pilots focused on bachata, salsa, merengue, cumbia, son/bolero/guaracha, and tropical fusion. We prefer full-length lossless WAV, stems, lyrics, MIDI where available, rich metadata, singing-voice consent, commercial model deployment, and permission to host trained weights/adapters. Please review the attached RFI and advise availability, pricing, sample access, and license terms."),
    ("SIGNAL.DAT", "Contact Curation Team", "Custom tropical Latin catalog and commercial training license", "Gen Audius LLC is assessing a tropical Latin pilot dataset for QAMUZ, our flagship music-generation brand. We are interested in your commercial training catalog and custom curation, particularly bachata, salsa, merengue, cumbia and related styles. Please confirm exact track/hour coverage, whether audio and stems are delivered or the product is metadata-only, rights provenance, voice-consent coverage, and pricing for 50-, 100-, and 300-hour equivalents. The attached RFI contains our technical and licensing requirements."),
    ("Sonovox", "Official licensing contact", "Spanish singing-voice dataset for tropical Latin music", "We are seeking licensed Spanish singing-voice data for a commercial tropical Latin music model. Our priorities are male and female leads, backing choruses, duet-capable performances, dry/wet stems, full-length phrases or songs, lyrics/timing, Caribbean Spanish coverage, and explicit AI-training consent. Please provide catalog size, voice distribution, perpetual-license terms, restrictions on identity imitation, pricing, and an evaluation sample. The attached RFI provides the broader project context."),
    ("LANDR Fair Trade AI", "Business / AI partnerships", "Developer access to rights-verified tropical Latin catalog", "Gen Audius LLC would like to explore licensing a custom subset of LANDR's opt-in Fair Trade AI catalog for QAMUZ, our flagship music-generation brand. We seek tropical Latin recordings, especially bachata, salsa, merengue and cumbia, with clear recording/publishing rights and permission for commercial AI training and outputs. Please advise whether developer dataset licensing is currently available, approximate eligible catalog coverage, term/perpetual options, pricing and the appropriate partnership contact."),
]
for vendor, contact, subject, body in messages:
    doc.add_heading(vendor, level=2)
    add_label_para(doc, "Contact", contact)
    add_label_para(doc, "Subject", subject)
    p = doc.add_paragraph(body)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_after = Pt(8)

doc.add_heading("10. Resumen interno (español)", level=1)
p = doc.add_paragraph("Este documento solicita cotizaciones comparables a nombre de Gen Audius LLC sin revelar un presupuesto. QAMUZ se presenta como la marca y el producto principal de la compañía. La prioridad es obtener derechos claros para entrenar, desplegar y comercializar QAMUZ, incluyendo los pesos resultantes. Antes de firmar, Gen Audius LLC debe revisar el contrato con asesoría jurídica especializada en propiedad intelectual y derechos musicales.")
p.paragraph_format.space_after = Pt(6)
for text in [
    "No aceptar una licencia de sincronizacion normal como sustituto de una licencia explicita para entrenamiento de IA.",
    "Confirmar por separado máster, composición, letra, interpretación, voz/publicidad y sublicencias.",
    "Comenzar con una muestra de evaluación y un piloto pequeño antes de adquirir 300 horas.",
    "Mantener el dataset tropical separado del futuro corpus MIDI multigénero.",
]:
    add_bullet(doc, text)

doc.add_heading("11. Reference links", level=1)
refs = [
    "GCX datasets and contact: https://www.gcx.co/",
    "GCX pricing approach: https://www.gcx.co/pricing",
    "SIGNAL.DAT licensing and pricing: https://signaldat.com/",
    "Sonovox singing-voice datasets: https://sonovox.ai/",
    "LANDR Fair Trade AI: https://www.landr.com/fairai",
    "Google Groove MIDI Dataset (CC BY 4.0): https://magenta.tensorflow.org/datasets/groove",
    "Slakh2100 (CC BY 4.0): https://www.slakh.com/",
    "UMD-600MB restricted MIDI dataset: https://zenodo.org/records/14560336",
]
for ref in refs:
    add_bullet(doc, ref)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(0)
r = p.add_run("Note: This RFI is a procurement and technical-scoping document, not legal advice. Final agreements require legal review.")
set_font(r, size=9, italic=True, color=GRAY)

doc.core_properties.title = "Gen Audius LLC / QAMUZ Licensed Music Dataset RFI"
doc.core_properties.subject = "Tropical Latin music and singing-voice dataset procurement"
doc.core_properties.author = "Gen Audius LLC"
doc.core_properties.keywords = "Gen Audius LLC, QAMUZ, dataset licensing, tropical Latin, AI music"
doc.save(OUT)
print(OUT)
